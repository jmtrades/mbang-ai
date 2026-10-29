"""One-click webinar launch: one brief in, the whole webinar built, scheduled and waiting on ONE approval.

A launch runs the webinar chain (offer, hooks, slide script, pages, show-up and replay sequences, run-of-show,
chat prompts, sales hand-off), creates and schedules the event, and queues a single "launch pack". Approving the
pack approves the registration page (so /w/<tenant>/<slug> goes live) and every show-up / replay email template
(so reminders fan out to each registrant against the start time, exactly as approving them one by one does).
The slide script is also available as a real .pptx deck.

Nothing new is sent by approving the pack: emails still go out only to people who register and consent, each at
its scheduled time. A student may launch for themselves or for one of their own clients, and sees only their own
launches.
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from . import brief as brief_mod
from . import config, db, queue, webinar_chain, webinars

KIND = "webinar_launch"
DEFAULT_TZ = "America/New_York"
DEFAULT_HOUR = 19


# A growth operator's client sells to consumers (traders, gym-goers, singles...), so the webinar speaks to THOSE
# people. When the client's brief has no audience/outcome yet, these fill the gaps by niche keyword. Explicit brief
# fields always win; B2B niches (agency, coaching businesses) keep the business-growth defaults.
AUDIENCE_PRESETS: list[tuple[tuple[str, ...], dict]] = [
    (("trad", "forex", "crypto", "stock", "futures", "options", "day trad"),
     {"audience": "new traders", "outcome": "consistent results", "mechanism": "3-step trade plan",
      "pain": "blowing up their account", "timeframe": "90 days"}),
    (("fitness", "gym", "fat loss", "weight", "personal train", "nutrition"),
     {"audience": "busy professionals", "outcome": "lean and strong", "mechanism": "3-day training plan",
      "pain": "giving up the food they love", "timeframe": "12 weeks"}),
    (("dating", "relationship", "attraction"),
     {"audience": "single men", "outcome": "a calendar full of great dates", "mechanism": "attraction framework",
      "pain": "swiping endlessly", "timeframe": "60 days"}),
    (("real estate", "property", "airbnb", "rental"),
     {"audience": "first-time investors", "outcome": "their first cash-flowing deal",
      "mechanism": "deal-finding system", "pain": "draining their savings", "timeframe": "6 months"}),
    (("ecom", "e-com", "amazon", "shopify", "dropship"),
     {"audience": "new store owners", "outcome": "their first profitable product",
      "mechanism": "product-testing method", "pain": "burning cash on ads", "timeframe": "90 days"}),
]
_PRESET_FIELDS = ("audience", "outcome", "mechanism", "pain", "timeframe")


def audience_preset(niche: str) -> dict:
    low = (niche or "").lower()
    for keys, preset in AUDIENCE_PRESETS:
        if any(k in low for k in keys):
            return preset
    return {}


def _fill_audience(partial: dict, for_client: bool) -> dict:
    """Fill the audience-side fields the brief left empty: by niche preset, else (for a client) a neutral
    consumer-facing fallback so the copy never reads as 'business owners get more enquiries'."""
    preset = audience_preset(str(partial.get("niche") or "")) or audience_preset(str(partial.get("offer") or ""))
    if not preset and for_client:
        niche = str(partial.get("niche") or "this").strip()
        preset = {"audience": f"people serious about {niche}", "outcome": "real, measurable progress",
                  "mechanism": "step-by-step method", "pain": "wasting months on trial and error",
                  "timeframe": "90 days"}
    return {**{k: preset[k] for k in _PRESET_FIELDS if k in preset and not partial.get(k)}, **partial}


class LaunchError(Exception):
    def __init__(self, status: int, msg: str):
        super().__init__(msg)
        self.status = status


def _actor(user: dict | None) -> str:
    return f"{user['username']}({user['role']})" if user else "system"


def _is_owner(user: dict | None) -> bool:
    return not user or user.get("role") == "owner"


def _client(tenant, client_id: int, user: dict | None) -> dict:
    with tenant.db() as c:
        found = db.rows(c.execute("SELECT * FROM clients WHERE id=?", (int(client_id),)))
    row = found[0] if found else None
    if not row or (not _is_owner(user) and row.get("owner_user_id") not in (None, (user or {}).get("id"))):
        raise LaunchError(404, "Client not found")
    return row


def launch(tenant, data: dict, user: dict | None = None) -> dict:
    """Build, schedule and queue a whole webinar from one brief. Returns the launch summary."""
    if not data.get("starts_at"):
        raise LaunchError(422, "When is the webinar? Give a date and time.")
    tz = str(data.get("tz") or DEFAULT_TZ)
    partial: dict = {}
    client = None
    if data.get("client_id"):
        client = _client(tenant, int(data["client_id"]), user)
        cb = brief_mod.for_client(client)
        # only what the client actually told us: build() back-fills business-growth defaults we must not keep
        partial.update({k: v for k, v in cb["brief"].items() if v not in (None, "") and k not in cb["missing"]})
    partial.update({k: data[k] for k in brief_mod.FIELDS if data.get(k) not in (None, "")})
    partial = _fill_audience(partial, for_client=client is not None)
    try:
        starts = webinars.parse_start(data["starts_at"], tz)
    except webinars.WebinarError as e:
        raise LaunchError(e.status, str(e)) from None
    partial["date"] = webinars.when_label(starts, tz)  # emails and pages name the real date, not a placeholder
    chain = webinar_chain.run_chain(tenant, partial, user)
    try:
        w = webinars.create(tenant, {"chain_id": chain["chain_id"], "starts_at": starts, "tz": tz,
                                     "title": data.get("title") or "", "client_id": client["id"] if client else None,
                                     "duration_min": data.get("duration_min") or 60}, user)
    except webinars.WebinarError as e:
        raise LaunchError(e.status, str(e)) from None
    web = webinars.get(tenant, w["id"])
    if web is None:
        raise LaunchError(500, "The webinar was not saved")
    when = webinars.when_label(web["starts_at"], web["tz"])
    n_tpl = len(w["template_queue_ids"])
    minutes = sum(m for _, m in webinar_chain.SEGMENTS)
    body = "\n".join([
        f"Approve once to launch \"{web['title']}\" on {when}.",
        "",
        "Approving this will:",
        "- Publish the registration page so people can sign up.",
        f"- Turn on {n_tpl} show-up and replay emails. Each registrant gets them at the right time before and after "
        "the event, and only if they agreed to emails.",
        f"- Keep the slide deck ready to download ({len(webinar_chain.SEGMENTS)} segments, about {minutes} minutes).",
        "",
        "Already drafted: offer, title and hooks, slide script, registration and thank-you pages, show-up and replay "
        "sequences, run-of-show, chat and offer-moment prompts, and the hand-off to sales.",
    ])
    meta = {"webinar_id": w["id"], "chain_id": chain["chain_id"], "page_qid": w["page_queue_id"],
            "template_qids": w["template_queue_ids"], "chain_qids": chain["queue_ids"],
            "client_id": client["id"] if client else None}
    if user and not _is_owner(user):
        meta["owner_user_id"] = user["id"]
    pack = queue.add(tenant, KIND, body, channel="manual", title=f"Launch pack: {web['title']}"[:200], meta=meta,
                     run_id=chain.get("run_id"), actor=_actor(user), demo=tenant.is_demo)
    return summary(tenant, w["id"], user, pack=pack, chain=chain)


def _pack_for(tenant, wid: int) -> dict | None:
    with tenant.db() as c:
        r = c.execute("SELECT id FROM queue WHERE kind=? AND json_extract(meta,'$.webinar_id')=? "
                      "ORDER BY id DESC LIMIT 1", (KIND, wid)).fetchone()
    return queue.get(tenant, r["id"]) if r else None


def _webinar(tenant, wid: int, user: dict | None) -> dict:
    w = webinars.get(tenant, wid)
    mine = _is_owner(user) or w is not None and w.get("owner_user_id") == (user or {}).get("id")
    if not w or not _pack_for(tenant, wid) or not mine:
        raise LaunchError(404, "Launch not found")
    return w


def summary(tenant, wid: int, user: dict | None = None, pack: dict | None = None, chain: dict | None = None) -> dict:
    w = _webinar(tenant, wid, user)
    pack = pack or _pack_for(tenant, wid)
    try:
        st = webinar_chain.state(tenant, w["chain_id"])
        assets = [{"step": s["step"], "label": s["label"], "body": ((s.get("output") or {}).get("body") or "")}
                  for s in st["steps"]]
    except KeyError:
        assets = []
    live = w.get("status") in ("open", "live", "done") and bool(w.get("page"))
    with tenant.db() as c:
        tpl = db.rows(c.execute("SELECT key,status FROM webinar_templates WHERE webinar_id=?", (wid,)))
        n_reg = c.execute("SELECT COUNT(*) FROM webinar_registrants WHERE webinar_id=?", (wid,)).fetchone()[0]
    return {"webinar_id": wid, "title": w["title"], "when": webinars.when_label(w["starts_at"], w["tz"]),
            "starts_at": w["starts_at"], "tz": w["tz"], "status": w.get("status"), "live": live,
            "page_url": webinars.page_url(tenant, w), "deck_url": f"/t/{tenant.id}/api/launches/{wid}/deck.pptx",
            "pack": {"id": pack["id"], "status": pack["status"], "body": pack["body"]} if pack else None,
            "emails": {"total": len(tpl), "on": sum(1 for t in tpl if t["status"] == "approved")},
            "registrants": n_reg, "assets": assets, "chain_id": w["chain_id"],
            "run_id": (chain or {}).get("run_id")}


def list_launches(tenant, user: dict | None = None) -> list[dict]:
    with tenant.db() as c:
        ids = [r[0] for r in c.execute("SELECT DISTINCT json_extract(meta,'$.webinar_id') FROM queue WHERE kind=? "
                                       "ORDER BY id DESC", (KIND,))]
    out = []
    for wid in ids:
        try:
            s = summary(tenant, int(wid), user)
        except LaunchError:
            continue
        s.pop("assets", None)
        out.append(s)
    return out


@queue.on_approve(KIND)
def _pack_approved(tenant, item: dict, actor: str) -> None:
    """Approve the page and every email template the pack stands for. Their own approve hooks do the work: the page
    goes live and the templates fan out per registrant. Guardrails still run on each one."""
    m = item["meta"]
    problems = []
    for qid in [m.get("page_qid"), *(m.get("template_qids") or [])]:
        if not qid:
            continue
        it = queue.get(tenant, int(qid), scrub=False)
        if not it or it["status"] not in queue.OPEN:
            continue
        try:
            queue.approve(tenant, int(qid), actor, role="owner")
        except queue.QueueError as e:
            problems.append(f"{it['title']}: {e}")
    if problems:
        raise queue.HookFailed("Launched, but these need a fix in Approvals: " + "; ".join(problems)[:400], status=409)


# --------------------------------------------------------------------------- slide deck
def _host(tenant, w: dict) -> str:
    """Whose webinar it is: the client's name when launched for a client, else this workspace's brand."""
    if w.get("client_id"):
        with tenant.db() as c:
            found = db.rows(c.execute("SELECT name FROM clients WHERE id=?", (int(w["client_id"]),)))
        if found and found[0].get("name"):
            return str(found[0]["name"])
    return config.get().brand


_SMALL = {"a", "an", "and", "as", "at", "but", "by", "for", "in", "of", "on", "or", "the", "to", "with"}


def headline(text: str) -> str:
    """Headline case for covers and pages: capitalise each word except short joining words (never the first)."""
    words = text.split()
    return " ".join(w if (i and w.lower() in _SMALL) else (w[:1].upper() + w[1:]) for i, w in enumerate(words))


# deck palette: deep ink for statement slides, warm accent, soft paper for teaching slides
_INK, _ACCENT, _PAPER, _TEXT, _MUTED, _PANEL = "0F172A", "F59E0B", "F8FAFC", "0F172A", "64748B", "1E293B"
_EMU_IN = 914400


def _rgb(h: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(h)


def _in(x: float) -> int:
    return int(x * _EMU_IN)


class _Deck:
    """A 16:9 deck painted on blank slides: statement slides in ink, teaching slides on paper, one accent."""

    def __init__(self, host: str):
        from pptx import Presentation

        self.prs = Presentation()
        from pptx.util import Emu

        self.prs.slide_width, self.prs.slide_height = Emu(_in(13.333)), Emu(_in(7.5))
        self.W, self.H = _in(13.333), _in(7.5)
        self.host, self.n = host, 0

    def rect(self, s, x: int, y: int, cx: int, cy: int, color: str, rounded: bool = False):
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE

        r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, cx, cy)
        r.fill.solid()
        r.fill.fore_color.rgb = RGBColor.from_string(color)
        r.line.fill.background()
        r.shadow.inherit = False
        return r

    def dot(self, s, x: int, y: int, size: int, color: str):
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE

        r = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
        r.fill.solid()
        r.fill.fore_color.rgb = RGBColor.from_string(color)
        r.line.fill.background()

    def text(self, s, box: tuple[float, float, float, float], value: str | list[str], size: int, color: str,
             bold: bool = False, align: str = "left", spacing: float = 1.1):
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Pt

        x, y, cx, cy = box
        tf = s.shapes.add_textbox(_in(x), _in(y), _in(cx), _in(cy)).text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, ln in enumerate(value if isinstance(value, list) else [value]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = spacing
            run = p.add_run()
            run.text = ln
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, "Calibri"
            run.font.color.rgb = RGBColor.from_string(color)

    def slide(self, dark: bool, note: str = ""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        self.rect(s, 0, 0, self.W, self.H, _INK if dark else _PAPER)
        if not dark:
            self.rect(s, 0, 0, _in(0.18), self.H, _ACCENT)
        if note:
            s.notes_slide.notes_text_frame.text = note
        return s

    def footer(self, s, dark: bool) -> None:
        col = "94A3B8" if dark else _MUTED
        self.text(s, (0.8, 6.9, 8, 0.3), self.host.upper(), 11, col, bold=True)
        self.text(s, (11.53, 6.9, 1, 0.3), str(self.n), 11, col, align="right")

    def kicker(self, s, label: str, y: float = 0.8) -> None:
        self.rect(s, _in(0.8), _in(y + 0.08), _in(0.5), 45720, _ACCENT)
        self.text(s, (1.45, y, 8, 0.35), label.upper(), 13, _ACCENT, bold=True)

    def bullets(self, title: str, label: str, items: list[str], note: str = "") -> None:
        s = self.slide(False, note)
        self.kicker(s, label)
        self.text(s, (0.8, 1.3, 11.5, 1.1), title, 38, _TEXT, bold=True)
        y = 2.7
        for it in [i for i in items if i][:6]:
            self.dot(s, _in(0.85), _in(y + 0.13), _in(0.16), _ACCENT)
            self.text(s, (1.3, y, 10.8, 0.7), it, 22, _TEXT)
            y += 0.78
        self.footer(s, False)

    def save(self) -> bytes:
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


def _deck_cover(d: _Deck, title: str, when: str, usp: str) -> None:
    s = d.slide(True, f"Welcome everyone in. Say who this is for and what they'll walk away with. {usp}".strip())
    d.rect(s, d.W - _in(4.2), 0, _in(4.2), d.H, _PANEL)
    d.rect(s, d.W - _in(4.2), 0, _in(0.12), d.H, _ACCENT)
    d.kicker(s, "Free live training", 1.2)
    d.text(s, (0.8, 1.8, 8.0, 3.2), headline(title), 44, "FFFFFF", bold=True, spacing=1.0)
    d.text(s, (0.8, 5.3, 8.0, 0.5), when, 20, _ACCENT, bold=True)
    d.text(s, (0.8, 5.85, 8.0, 0.5), f"Hosted by {d.host}", 16, "CBD5E1")
    total = sum(m for _, m in webinar_chain.SEGMENTS)
    d.text(s, (9.63, 2.2, 3.2, 1.4), str(total), 88, "FFFFFF", bold=True, align="center", spacing=0.9)
    d.text(s, (9.63, 3.6, 3.2, 0.5), "MINUTES, LIVE", 14, _ACCENT, bold=True, align="center")
    d.text(s, (9.63, 4.4, 3.2, 1.0), ["Stay to the end", "for live Q&A"], 18, "CBD5E1", align="center")


def _deck_agenda(d: _Deck) -> None:
    s = d.slide(False, "Walk the agenda in 20 seconds so people know to stay for the offer and the Q&A.")
    d.kicker(s, "Agenda")
    d.text(s, (0.8, 1.3, 11.5, 1.0), "Here's how the next hour goes", 38, _TEXT, bold=True)
    for i, (name, mins) in enumerate(webinar_chain.SEGMENTS):
        x, y = 0.8 + 2.95 * (i % 4), 2.65 + 1.95 * (i // 4)
        tile = d.rect(s, _in(x), _in(y), _in(2.75), _in(1.75), "FFFFFF", rounded=True)
        tile.line.color.rgb = _rgb("E2E8F0")
        d.text(s, (x + 0.25, y + 0.2, 1.0, 0.5), f"{i + 1:02d}", 22, _ACCENT, bold=True)
        d.text(s, (x + 0.25, y + 0.7, 2.25, 0.6), name, 17, _TEXT, bold=True)
        d.text(s, (x + 0.25, y + 1.25, 2.25, 0.35), f"{mins} min", 13, _MUTED)
    d.footer(s, False)


def _deck_segments(d: _Deck, notes: dict[int, str]) -> None:
    for i, (name, mins) in enumerate(webinar_chain.SEGMENTS, start=1):
        point = re.sub(r"^[A-Za-z ]+:\s*", "", notes.get(i, "")).strip().rstrip(".")
        dark = i in (1, 4)  # the promise and the mechanism get the statement treatment
        s = d.slide(dark, f"Segment {i}: {name}, {mins} min. {notes.get(i, '')}".strip())
        d.text(s, (0.8, 0.9, 3, 1.6), f"{i:02d}", 96, _ACCENT, bold=True, spacing=0.9)
        d.kicker(s, f"Part {i} · {mins} min", 2.7)
        d.text(s, (0.8, 3.2, 11.5, 1.0), name, 44, "FFFFFF" if dark else _TEXT, bold=True)
        if point:
            d.text(s, (0.8, 4.4, 11.0, 1.6), point[:1].upper() + point[1:], 24, "CBD5E1" if dark else "334155")
        d.footer(s, dark)


def _deck_close(d: _Deck, offer: str, usp: str, price: float) -> None:
    s = d.slide(True, "Who it is for and who it is not for. Only verified results, with permission.")
    d.kicker(s, "The offer")
    d.text(s, (0.8, 1.4, 7.6, 1.6), offer[:1].upper() + offer[1:], 44, "FFFFFF", bold=True)
    d.text(s, (0.8, 3.2, 7.4, 2.6), usp, 22, "CBD5E1")
    card = d.rect(s, _in(9.0), _in(1.6), _in(3.5), _in(3.6), _PANEL, rounded=True)
    card.line.color.rgb = _rgb(_ACCENT)
    d.text(s, (9.0, 2.1, 3.5, 0.4), "INVESTMENT", 13, _ACCENT, bold=True, align="center")
    d.text(s, (9.0, 2.6, 3.5, 1.0), f"${price:,.0f}" if price else "Ask on the call", 40 if price else 24, "FFFFFF",
           bold=True, align="center")
    d.text(s, (9.2, 3.8, 3.1, 1.0), "Book a call to see if it's a fit", 16, "CBD5E1", align="center")
    d.footer(s, True)
    s = d.slide(True, "Stay on for live questions. Drop the booking link in the chat twice.")
    d.rect(s, 0, d.H - _in(1.2), d.W, _in(1.2), _PANEL)
    d.kicker(s, "Next step", 2.0)
    d.text(s, (0.8, 2.6, 11.5, 1.4), "Book your call", 60, "FFFFFF", bold=True)
    d.text(s, (0.8, 4.1, 11.0, 1.0), "We'll map out your plan together and see if it's a fit. Questions? Ask in the "
           "chat now.", 22, "CBD5E1")
    d.footer(s, True)


def deck(tenant, wid: int, user: dict | None = None) -> tuple[str, bytes]:
    """The slide script as a designed 16:9 PowerPoint: cover, promise, agenda, one slide per segment, the offer and
    the next step, each with speaker notes."""
    w = _webinar(tenant, wid, user)
    st = webinar_chain.state(tenant, w["chain_id"])
    b = st.get("brief") or {}
    outs = {s["step"]: (s.get("output") or {}) for s in st["steps"]}
    hooks = ((outs.get("title_hooks") or {}).get("data") or {}).get("hooks") or []
    usp = ((outs.get("offer") or {}).get("data") or {}).get("usp") or ""
    notes = {}
    for line in ((outs.get("slide_script") or {}).get("body") or "").splitlines():
        mm = re.match(r"Segment (\d+): .*?—\s*(.+)$", line)
        if mm:
            notes[int(mm.group(1))] = mm.group(2).strip()
    d = _Deck(_host(tenant, w))
    _deck_cover(d, w["title"], webinars.when_label(w["starts_at"], w["tz"]), usp)
    d.bullets("What you'll learn today", "The promise", hooks[:4] or [usp],
              "Open with the promise. Keep it under a minute, then ask them to type where they are from.")
    _deck_agenda(d)
    _deck_segments(d, notes)
    _deck_close(d, str(b.get("offer") or "The programme"), usp, float(b.get("price") or 0))
    name = re.sub(r"[^A-Za-z0-9]+", "-", w["title"]).strip("-")[:60] or "webinar"
    return f"{name}.pptx", d.save()


# --------------------------------------------------------------------------- chat: "launch a webinar ... on Nov 12"
LAUNCH_ASK = re.compile(r"\b(launch|run|set up|schedule|host)\b[^.?!\n]{0,40}\bwebinar\b", re.I)
_MON = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
_MONTHS = {m: i for i, m in enumerate(_MON, 1)}
_DATE = re.compile(r"\b(?P<mon>" + "|".join(_MON) + r")[a-z]*\.?\s+(?P<day>\d{1,2})(?:st|nd|rd|th)?\b"
                   r"|\b(?P<iso>\d{4}-\d{2}-\d{2})\b", re.I)
_TIME = re.compile(r"\b(?:at\s+)?(?P<h>\d{1,2})(?::(?P<m>\d{2}))?\s*(?P<ap>am|pm)\b", re.I)
_TZ = {"pt": "America/Los_Angeles", "pst": "America/Los_Angeles", "pdt": "America/Los_Angeles",
       "pacific": "America/Los_Angeles", "et": "America/New_York", "est": "America/New_York",
       "edt": "America/New_York", "eastern": "America/New_York",
       "ct": "America/Chicago", "cst": "America/Chicago", "central": "America/Chicago", "uk": "Europe/London",
       "gmt": "Europe/London", "bst": "Europe/London", "london": "Europe/London"}
_CLIENT = re.compile(r"\bfor (?:my |our )?(?P<what>[\w &'-]{2,60}?) client\b", re.I)


def wants_launch(text: str) -> bool:
    """A launch ask WITH a date runs the one-click launch; without one it stays a mission goal (its intake asks)."""
    t = text or ""
    return bool(LAUNCH_ASK.search(t)) and bool(_DATE.search(t) or re.search(r"\bnext week\b", t, re.I))


def parse_request(text: str, now: datetime | None = None) -> dict:
    """Pull the date, time, time zone and client hint out of a chat message. Missing date -> no starts_at."""
    tz = DEFAULT_TZ
    for word in re.findall(r"[a-z]+", (text or "").lower()):
        if word in _TZ:
            tz = _TZ[word]
            break
    try:
        zone = ZoneInfo(tz)
    except ZoneInfoNotFoundError:
        zone, tz = ZoneInfo("UTC"), "UTC"
    now = now or datetime.now(zone)
    out: dict = {"tz": tz}
    d = _DATE.search(text or "")
    if d:
        if d.group("iso"):
            day = datetime.fromisoformat(d.group("iso")).date()
        else:
            mon, dd = _MONTHS[d.group("mon").lower()[:3]], int(d.group("day"))
            try:
                day = now.replace(month=mon, day=dd).date()
            except ValueError:
                day = None
            if day and day < now.date():
                day = day.replace(year=day.year + 1)
        if day:
            hh, mi = DEFAULT_HOUR, 0
            tm = _TIME.search(text[d.end():] if d else text)
            if tm:
                hh = int(tm.group("h")) % 12 + (12 if tm.group("ap").lower() == "pm" else 0)
                mi = int(tm.group("m") or 0)
            out["starts_at"] = datetime(day.year, day.month, day.day, hh, mi).isoformat(timespec="minutes")
    elif re.search(r"\bnext week\b", text or "", re.I):
        day = (now + timedelta(days=7)).date()
        out["starts_at"] = datetime(day.year, day.month, day.day, DEFAULT_HOUR, 0).isoformat(timespec="minutes")
    cm = _CLIENT.search(text or "")
    if cm:
        out["client_hint"] = cm.group("what").strip()
    return out


def match_client(tenant, hint: str, user: dict | None) -> dict | None:
    if not hint:
        return None
    with tenant.db() as c:
        rows = db.rows(c.execute("SELECT * FROM clients ORDER BY id DESC"))
    if not _is_owner(user):
        rows = [r for r in rows if r.get("owner_user_id") == (user or {}).get("id")]
    words = [w for w in re.findall(r"[a-z]+", hint.lower()) if len(w) > 2]
    best, score = None, 0
    for r in rows:
        hay = f"{r.get('name', '')} {r.get('niche', '')} {r.get('offer', '')}".lower()
        s = sum(1 for w in words if w in hay)
        if s > score:
            best, score = r, s
    return best


def from_chat(tenant, text: str, user: dict) -> tuple[str, dict]:
    """Chat entry point. Returns (reply, meta)."""
    req = parse_request(text)
    if not req.get("starts_at"):
        return ("I can launch the whole webinar in one go: offer, hooks, slide deck, registration page and every "
                "reminder email. When is it? For example: 'Launch a webinar for my trading client on November 12 at "
                "7pm ET'."), {"launch": None}
    data = {"starts_at": req["starts_at"], "tz": req["tz"]}
    client = match_client(tenant, req.get("client_hint", ""), user)
    if client:
        data["client_id"] = client["id"]
    elif req.get("client_hint"):
        data["niche"] = req["client_hint"]
    try:
        res = launch(tenant, data, user)
    except LaunchError as e:
        return str(e), {"launch": None}
    who = f" for {client['name']}" if client else ""
    n = res["emails"]["total"]
    reply = (f"Done. I built your webinar{who}: \"{res['title']}\" on {res['when']}. Offer, hooks, a "
             f"{len(webinar_chain.SEGMENTS)}-part slide deck, registration and thank-you pages, {n} "
             "show-up and replay emails, run-of-show, chat prompts and the sales hand-off are ready. Approve the "
             "launch pack once and the page goes live and the reminders switch on.")
    return reply, {"launch": res}
